import io
from typing import Optional

import openpyxl
import pandas as pd

from app.csv_io.parser import MAX_ROWS, FileParsingError, build_flux_dataframe, sanitize_cell_for_export
from app.dataframe.frame import FluxDataFrame
from app.dataframe.schema import DataSchema


def list_sheet_names(data: bytes) -> list[str]:
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
    except Exception as e:
        raise FileParsingError(f"Could not read XLSX file: {e}") from e
    return workbook.sheetnames


def parse_xlsx(data: bytes, sheet_name: Optional[str] = None, max_rows: int = MAX_ROWS) -> FluxDataFrame:
    """Parse XLSX bytes into a FluxDataFrame. Defaults to the first sheet if sheet_name is blank.

    Runs the same inference pass as CSV (via build_flux_dataframe) — Excel's own cell typing
    (openpyxl hands back native int/float/datetime/bool objects) doesn't map 1:1 onto our
    DataType/pandas dtypes, so every cell is stringified first rather than trusted blindly.
    """
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as e:
        raise FileParsingError(f"Could not read XLSX file: {e}") from e

    if sheet_name:
        if sheet_name not in workbook.sheetnames:
            raise FileParsingError(
                f"Sheet '{sheet_name}' not found. Available sheets: {', '.join(workbook.sheetnames)}"
            )
        sheet = workbook[sheet_name]
    else:
        sheet = workbook.worksheets[0]

    rows_iter = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return FluxDataFrame(schema=DataSchema(columns=[]), df=pd.DataFrame(), total_row_count=0)

    headers = [
        str(h) if h is not None and str(h).strip() != "" else f"Column_{i}"
        for i, h in enumerate(header_row)
    ]

    data_rows = []
    for row in rows_iter:
        if row == (None,) * len(row):
            continue  
        data_rows.append([_stringify_cell(v) for v in row])
        if len(data_rows) > max_rows:
            raise FileParsingError(
                f"File has more than {max_rows:,} rows, which exceeds the upload limit. "
                "Trim the file or split it before uploading."
            )

    raw_df = pd.DataFrame(data_rows, columns=headers)
    return build_flux_dataframe(raw_df)


def _stringify_cell(value) -> Optional[str]:
    if value is None:
        return None
    return str(value)


def write_xlsx(df: pd.DataFrame, sheet_name: str = "Sheet1") -> bytes:
    """Serialize a DataFrame to XLSX bytes, sanitized against formula injection on export."""
    sanitized = df.map(sanitize_cell_for_export)
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        sanitized.to_excel(writer, sheet_name=sheet_name, index=False)
    return buffer.getvalue()
