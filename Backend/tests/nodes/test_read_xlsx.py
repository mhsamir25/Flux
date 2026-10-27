import io

import openpyxl
import pytest

from app.nodes.base import FieldType, NodeCategory, SchemaValidationError
from app.nodes.source.read_xlsx import ReadXlsxNode


def _make_xlsx(rows: list[list], sheet_name: str = "Sheet1") -> bytes:
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = sheet_name
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def test_id_and_category():
    node = ReadXlsxNode()
    assert node.id == "read_xlsx"
    assert node.category == NodeCategory.SOURCE
    assert node.input_ports == 0
    assert node.output_ports == 1


def test_config_schema_has_optional_sheet_name_field():
    node = ReadXlsxNode()
    fields = node.get_config_schema()
    assert len(fields) == 1
    assert fields[0].key == "sheet_name"
    assert fields[0].type == FieldType.TEXT
    assert fields[0].required is False


def test_infer_output_schema_is_empty_placeholder():
    node = ReadXlsxNode()
    schema = node.infer_output_schema(None, {})
    assert schema.columns == []


def test_execute_raises_not_implemented():
    node = ReadXlsxNode()
    with pytest.raises(NotImplementedError):
        node.execute(None, {})


def test_execute_source_missing_data_raises_schema_validation_error():
    node = ReadXlsxNode()
    with pytest.raises(SchemaValidationError):
        node.execute_source({})


def test_execute_source_parses_default_sheet():
    data = _make_xlsx([["name", "age"], ["Alice", 30], ["Bob", 25]])
    node = ReadXlsxNode()
    fdf = node.execute_source({"_xlsx_bytes": data})
    assert fdf.total_row_count == 2
    assert fdf.schema.column_names() == ["name", "age"]


def test_execute_source_blank_sheet_name_defaults_to_first_sheet():
    data = _make_xlsx([["a"], [1]])
    node = ReadXlsxNode()
    fdf = node.execute_source({"_xlsx_bytes": data, "sheet_name": "   "})
    assert fdf.schema.column_names() == ["a"]


def test_execute_source_selects_named_sheet():
    wb = openpyxl.Workbook()
    wb.active.title = "First"
    wb.active.append(["a"])
    second = wb.create_sheet("Second")
    second.append(["b"])
    buffer = io.BytesIO()
    wb.save(buffer)
    data = buffer.getvalue()

    node = ReadXlsxNode()
    fdf = node.execute_source({"_xlsx_bytes": data, "sheet_name": "Second"})
    assert fdf.schema.column_names() == ["b"]


def test_execute_source_unknown_sheet_raises_schema_validation_error():
    data = _make_xlsx([["a"], [1]])
    node = ReadXlsxNode()
    with pytest.raises(SchemaValidationError):
        node.execute_source({"_xlsx_bytes": data, "sheet_name": "DoesNotExist"})


def test_to_excel_formula_placeholder():
    node = ReadXlsxNode()
    assert node.to_excel_formula(None, {}) == "Sheet1!A1:Z1000"


def test_to_excel_formula_explanation_mentions_adjusting_range():
    node = ReadXlsxNode()
    explanation = node.to_excel_formula_explanation({})
    assert "adjust" in explanation.lower()
