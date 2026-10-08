import io

import openpyxl
import pandas as pd
import pytest

from app.csv_io.parser import (
    FileParsingError,
    parse_csv,
    sanitize_cell_for_export,
    sniff_delimiter,
    write_csv,
)
from app.csv_io.xlsx_io import list_sheet_names, parse_xlsx, write_xlsx
from app.dataframe.schema import DataType


#CSV Parsing part

def test_parse_csv_basic():
    csv = "name,age,score\nAlice,30,95.5\nBob,25,87.3"
    fdf = parse_csv(csv)

    assert fdf.total_row_count == 2
    assert fdf.schema.column_names() == ["name", "age", "score"]
    assert fdf.schema.get_column("age").type == DataType.NUMBER
    assert fdf.df.iloc[0]["name"] == "Alice"
    assert fdf.df.iloc[0]["age"] == 30


def test_parse_csv_infers_date_and_string():
    csv = "event_date,product\n2026-01-15,Laptop\n2026-02-20,Mouse"
    fdf = parse_csv(csv)
    assert fdf.schema.get_column("event_date").type == DataType.DATE
    assert fdf.schema.get_column("product").type == DataType.STRING


def test_parse_csv_null_values():
    csv = "col1,col2,col3\n1,,value\n,2,\nvalue,,3"
    fdf = parse_csv(csv)
    assert pd.isna(fdf.df.iloc[0]["col2"])
    assert pd.isna(fdf.df.iloc[1]["col1"])


def test_parse_csv_empty_string_returns_empty_frame():
    fdf = parse_csv("")
    assert fdf.total_row_count == 0
    assert fdf.schema.columns == []


def test_sniff_delimiter_semicolon():
    sample = "name;age;score\nAlice;30;95.5\nBob;25;87.3"
    assert sniff_delimiter(sample) == ";"


def test_parse_csv_semicolon_delimited():
    csv = "name;age\nAlice;30\nBob;25"
    fdf = parse_csv(csv)
    assert fdf.schema.column_names() == ["name", "age"]
    assert fdf.df.iloc[0]["age"] == 30


def test_parse_csv_tab_delimited():
    csv = "name\tage\nAlice\t30\nBob\t25"
    fdf = parse_csv(csv)
    assert fdf.schema.column_names() == ["name", "age"]


def test_parse_csv_bytes_utf8():
    raw = "name,city\nAlice,Zürich".encode("utf-8")
    fdf = parse_csv(raw)
    assert fdf.df.iloc[0]["city"] == "Zürich"


def test_parse_csv_bytes_non_utf8_encoding_falls_back():
    raw = "name,city\nAlice,Zürich".encode("latin-1")
    fdf = parse_csv(raw)
    assert fdf.df.iloc[0]["name"] == "Alice"


def test_parse_csv_bom_stripped():
    raw = "﻿name,age\nAlice,30".encode("utf-8")
    fdf = parse_csv(raw)
    assert fdf.schema.column_names() == ["name", "age"]


def test_parse_csv_row_cap_exceeded_raises():
    header = "id,value\n"
    rows = "\n".join(f"{i},{i}" for i in range(10))
    csv = header + rows
    with pytest.raises(FileParsingError):
        parse_csv(csv, max_rows=5)


def test_parse_csv_boolean_inference():
    csv = "flag\ntrue\nfalse\nyes\nno"
    fdf = parse_csv(csv)
    assert fdf.schema.get_column("flag").type == DataType.BOOLEAN
    assert list(fdf.df["flag"]) == [True, False, True, False]



def test_sanitize_cell_for_export_prefixes_formula_chars():
    assert sanitize_cell_for_export("=SUM(A1)") == "'=SUM(A1)"
    assert sanitize_cell_for_export("+1+1") == "'+1+1"
    assert sanitize_cell_for_export("-cmd") == "'-cmd"
    assert sanitize_cell_for_export("@evil") == "'@evil"


def test_sanitize_cell_for_export_leaves_safe_values():
    assert sanitize_cell_for_export("hello") == "hello"
    assert sanitize_cell_for_export(42) == 42
    assert sanitize_cell_for_export(-5) == -5  # not a string, dtype-preserved negative number


def test_write_csv_sanitizes_injection_payload():
    df = pd.DataFrame({"note": ["=cmd|'/c calc'!A1", "safe text"], "amount": [-5, 10]})
    csv_text = write_csv(df)
    assert "'=cmd" in csv_text
    assert "-5" in csv_text  # numeric column untouched


# XLSX Parsing part

def _make_xlsx(rows: list[list], sheet_name: str = "Sheet1") -> bytes:
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = sheet_name
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def test_parse_xlsx_basic():
    data = _make_xlsx([["name", "age"], ["Alice", 30], ["Bob", 25]])
    fdf = parse_xlsx(data)
    assert fdf.total_row_count == 2
    assert fdf.schema.get_column("age").type == DataType.NUMBER
    assert fdf.df.iloc[0]["name"] == "Alice"


def test_parse_xlsx_missing_header_autonamed():
    data = _make_xlsx([[None, "name", None], [1, "Alice", 30]])
    fdf = parse_xlsx(data)
    names = fdf.schema.column_names()
    assert "Column_0" in names
    assert "name" in names


def test_parse_xlsx_empty_sheet():
    data = _make_xlsx([])
    fdf = parse_xlsx(data)
    assert fdf.total_row_count == 0
    assert fdf.schema.columns == []


def test_parse_xlsx_sheet_name_selection():
    wb = openpyxl.Workbook()
    wb.active.title = "First"
    wb.active.append(["a"])
    wb.active.append([1])
    second = wb.create_sheet("Second")
    second.append(["b"])
    second.append([2])
    buffer = io.BytesIO()
    wb.save(buffer)
    data = buffer.getvalue()

    fdf_default = parse_xlsx(data)
    assert fdf_default.schema.column_names() == ["a"]

    fdf_second = parse_xlsx(data, sheet_name="Second")
    assert fdf_second.schema.column_names() == ["b"]


def test_parse_xlsx_unknown_sheet_raises():
    data = _make_xlsx([["a"], [1]])
    with pytest.raises(FileParsingError):
        parse_xlsx(data, sheet_name="DoesNotExist")


def test_parse_xlsx_row_cap_exceeded_raises():
    rows = [["id"]] + [[i] for i in range(10)]
    data = _make_xlsx(rows)
    with pytest.raises(FileParsingError):
        parse_xlsx(data, max_rows=5)


def test_list_sheet_names():
    wb = openpyxl.Workbook()
    wb.active.title = "First"
    wb.create_sheet("Second")
    buffer = io.BytesIO()
    wb.save(buffer)
    assert list_sheet_names(buffer.getvalue()) == ["First", "Second"]


def test_parse_xlsx_invalid_bytes_raises():
    with pytest.raises(FileParsingError):
        parse_xlsx(b"not a real xlsx file")


def test_write_xlsx_sanitizes_injection_payload():
    df = pd.DataFrame({"note": ["=cmd|'/c calc'!A1", "safe"], "amount": [-5, 10]})
    xlsx_bytes = write_xlsx(df)

    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    sheet = wb.active
    values = [cell.value for cell in sheet["A"]]
    assert values[0] == "note"
    assert values[1] == "'=cmd|'/c calc'!A1"
    assert values[2] == "safe"
